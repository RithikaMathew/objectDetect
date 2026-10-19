#!/usr/bin/env python3
"""
aggregate_colonial_v11.py

Joins the per-hour outputs of xclip_overlay_v11_colonial.py with the Colonial North ground truth
(15-minute bins from the counter .xls) and writes, into <root>/validation/ :

    colonial_north_<run>_validation.xlsx   Per_15min | Per_hour | Metrics | Run_Log
    colonial_north_<run>_validation_log.txt

Nothing is overwritten: an existing file gets _v2, _v3, ... instead.

Mapping used (checked against the data):
  * ground truth .xls = SpreadsheetML XML, 192 bins = 48 h from 2019-07-10 00:00 (day 1 + day 2)
  * video chunk N (hourly mp4, .001 = 2019-07-10 00:00 on the camera's on-screen clock)
    covers GT bins (N-1)*4 .. (N-1)*4+3
  * bicyclist GT per bin = Bicycles sheet, bike lane (WB+EB) and sidewalk (WB+EB) of the NORTH channels
  * pedestrian GT per bin = Pedestrians sheet, first four (North) channels
"""
import argparse
import csv
import datetime as dt
import glob
import json
import math
import os
import re
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NS = {"ss": "urn:schemas-microsoft-com:office:spreadsheet"}
TIME_RE = re.compile(r"^\d{1,2}:\d{2} (AM|PM)$")


def unique_path(path):
    if not os.path.exists(path):
        return path
    root, ext = os.path.splitext(path)
    k = 2
    while os.path.exists(f"{root}_v{k}{ext}"):
        k += 1
    return f"{root}_v{k}{ext}"


# ----------------------------------------------------------------------------- ground truth
def read_sheet_rows(xls_path, sheet_name):
    root = ET.parse(xls_path).getroot()
    for ws in root.findall("ss:Worksheet", NS):
        if ws.get("{%s}Name" % NS["ss"]) == sheet_name:
            rows = []
            for row in ws.findall(".//ss:Row", NS):
                cells = []
                for c in row.findall("ss:Cell", NS):
                    d = c.find("ss:Data", NS)
                    cells.append((d.text or "") if d is not None else "")
                rows.append(cells)
            return rows
    raise SystemExit(f"ERROR: sheet {sheet_name!r} not found in {xls_path}")


def parse_gt(xls_path, ped_cols):
    """Returns dict bin_index -> {bic_bl, bic_sw, ped}, plus the 15-min labels and a check list."""
    checks = []
    bic_rows = read_sheet_rows(xls_path, "Bicycles")
    ped_rows = read_sheet_rows(xls_path, "Pedestrians")

    chan = next(r for r in bic_rows if r and r[0] == "Channel")[1:5]
    bl_cols = [i + 1 for i, c in enumerate(chan) if "bike lane" in c.lower()]
    sw_cols = [i + 1 for i, c in enumerate(chan) if "sidewalk" in c.lower()]
    checks.append(f"Bicycles channels used: {chan} -> bike-lane cols {bl_cols}, sidewalk cols {sw_cols}")

    bic = [r for r in bic_rows if r and TIME_RE.match(r[0])]
    ped = [r for r in ped_rows if r and TIME_RE.match(r[0])]
    if len(bic) != len(ped):
        checks.append(f"WARNING: Bicycles has {len(bic)} bins but Pedestrians has {len(ped)}")
    if len(bic) % 96:
        checks.append(f"WARNING: {len(bic)} bins is not a whole number of 24 h days")
    exp = [(dt.datetime(2000, 1, 1) + dt.timedelta(minutes=15 * (i % 96))).strftime("%I:%M %p").lstrip("0")
           for i in range(len(bic))]
    bad = [i for i, r in enumerate(bic) if r[0] != exp[i]]
    checks.append("Bin labels are consecutive 15-min steps starting 12:00 AM: " + ("YES" if not bad else f"NO, first bad index {bad[0]}"))

    gt = {}
    for i, r in enumerate(bic):
        g = lambda cols: sum(int(r[c] or 0) for c in cols)
        gt[i] = {"bic_bl": g(bl_cols), "bic_sw": g(sw_cols),
                 "ped": sum(int(ped[i][c] or 0) for c in ped_cols) if i < len(ped) else 0,
                 "label": r[0]}
    return gt, checks


def gt_csv_total(gt_csv, site):
    if not gt_csv or not os.path.exists(gt_csv):
        return None
    tot = 0
    with open(gt_csv, newline="") as f:
        for row in csv.DictReader(f):
            if row["site"] == site:
                tot += int(row["count"])
    return tot


def mapping_row(mapping_csv, site, date):
    if not mapping_csv or not os.path.exists(mapping_csv):
        return None
    with open(mapping_csv, newline="") as f:
        for row in csv.DictReader(f):
            if row["gt_site"] == site and row["date"] == date:
                return row
    return None


# ----------------------------------------------------------------------------- run outputs
def chunk_of(path):
    m = re.search(r"\.(\d{3})_(?:bin_counts|run_log)(?:_v\d+)?\.(?:csv|json)$", path)
    return int(m.group(1)) if m else None


def latest_per_chunk(pattern):
    best = {}
    for p in glob.glob(pattern):
        n = chunk_of(p)
        if n is not None and (n not in best or os.path.getmtime(p) > os.path.getmtime(best[n])):
            best[n] = p
    return best


def parse_chunks(spec):
    out = set()
    for part in spec.split(","):
        a, _, b = part.partition("-")
        out.update(range(int(a), int(b or a) + 1))
    return sorted(out)


# ----------------------------------------------------------------------------- python-side metrics (for the text log)
def py_metrics(ai, gt, inc):
    pairs = [(a, g) for a, g, i in zip(ai, gt, inc) if i]
    n = len(pairs)
    if n == 0:
        return {}
    err = [a - g for a, g in pairs]
    sa, sg = sum(a for a, _ in pairs), sum(g for _, g in pairs)
    matched = sum(min(a, g) for a, g in pairs)
    mx, my = sa / n, sg / n
    num = sum((a - mx) * (g - my) for a, g in pairs)
    den = math.sqrt(sum((a - mx) ** 2 for a, _ in pairs) * sum((g - my) ** 2 for _, g in pairs))
    return {
        "bins": n, "AI total": sa, "GT total": sg,
        "bias (AI-GT per bin)": sum(err) / n, "MAE": sum(abs(e) for e in err) / n,
        "RMSE": math.sqrt(sum(e * e for e in err) / n),
        "WAPE (sum|err|/sum GT)": (sum(abs(e) for e in err) / sg) if sg else float("nan"),
        "within +-1": sum(abs(e) <= 1 for e in err) / n, "within +-2": sum(abs(e) <= 2 for e in err) / n,
        "pearson r": (num / den) if den else float("nan"),
        "count-matched precision (approx)": matched / sa if sa else float("nan"),
        "count-matched recall (approx)": matched / sg if sg else float("nan"),
    }


# ----------------------------------------------------------------------------- workbook
ARIAL = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HDR_FILL = PatternFill("solid", start_color="1F3864")
HDR_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
INPUT_FONT = Font(name="Arial", size=10, color="0000FF")
INPUT_FILL = PatternFill("solid", start_color="FFFF00")


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook(rows, hours, runlog_rows, notes, args, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Per_15min"
    n = len(rows)
    last = n + 1
    heads = ["Video file", "Annotated clip", "Site", "Location / camera", "Date", "Bin start", "Bin end",
             "Chunk #", "Seconds processed", "Include (1/0)",
             "AI bicyclists (total)", "AI bicyclists bike lane", "AI bicyclists sidewalk", "AI bicyclists unassigned",
             "GT bicyclists bike lane", "GT bicyclists sidewalk", "GT bicyclists (total)",
             "Diff bicyclists (AI-GT)", "Abs diff", "Matched = MIN(AI,GT)", "Flag",
             "AI pedestrians", "GT pedestrians", "Diff pedestrians (AI-GT)", "AI vehicles"]
    ws.append(heads)
    style_header(ws, 1, len(heads))
    for i, r in enumerate(rows):
        x = i + 2
        ws.append([r["video"], r["clip"], args.site, args.location, r["date"], r["t0"], r["t1"], r["chunk"],
                   r["secs"], f"=IF(I{x}>=Metrics!$B$4*Metrics!$B$3,1,0)",
                   r["ai_bic"], r["ai_bl"], r["ai_sw"], r["ai_na"], r["gt_bl"], r["gt_sw"], f"=O{x}+P{x}",
                   f"=K{x}-Q{x}", f"=ABS(R{x})", f"=MIN(K{x},Q{x})",
                   f'=IF(J{x}=0,"partial",IF(S{x}>Metrics!$B$5,"CHECK","ok"))',
                   r["ai_ped"], r["gt_ped"], f"=V{x}-W{x}", r["ai_veh"]])
    for row in ws.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.font = ARIAL
    widths = [38, 44, 15, 30, 11, 10, 10, 8, 11, 9, 11, 11, 11, 11, 11, 11, 11, 12, 8, 11, 9, 11, 11, 12, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{last}"
    ws.conditional_formatting.add(f"A2:Y{last}", FormulaRule(formula=['$U2="CHECK"'], fill=PatternFill("solid", start_color="F8CBAD")))
    ws.conditional_formatting.add(f"A2:Y{last}", FormulaRule(formula=['$U2="partial"'], fill=PatternFill("solid", start_color="D9D9D9")))

    # ---- Per_hour
    wh = wb.create_sheet("Per_hour")
    hh = ["Chunk #", "Video file", "Hour start", "Bins included", "AI bicyclists", "GT bicyclists", "Diff (AI-GT)",
          "AI pedestrians", "GT pedestrians", "Diff pedestrians", "AI vehicles"]
    wh.append(hh)
    style_header(wh, 1, len(hh))
    R = lambda col: f"Per_15min!${col}$2:${col}${last}"
    for i, h in enumerate(hours):
        x = i + 2
        cond = f"({R('H')}=A{x})*{R('J')}"
        wh.append([h["chunk"], h["video"], h["t0"],
                   f"=SUMPRODUCT({cond})",
                   f"=SUMPRODUCT({cond}*{R('K')})", f"=SUMPRODUCT({cond}*{R('Q')})", f"=E{x}-F{x}",
                   f"=SUMPRODUCT({cond}*{R('V')})", f"=SUMPRODUCT({cond}*{R('W')})", f"=H{x}-I{x}",
                   f"=SUMPRODUCT({cond}*{R('Y')})"])
    tx = len(hours) + 2
    wh.append(["Total", "", "", f"=SUM(D2:D{tx - 1})"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{tx - 1})" for c in range(5, 12)])
    for row in wh.iter_rows(min_row=2, max_row=tx):
        for c in row:
            c.font = BOLD if row[0].row == tx else ARIAL
    for i, w in enumerate([9, 40, 18, 10, 12, 12, 12, 13, 13, 14, 11], 1):
        wh.column_dimensions[get_column_letter(i)].width = w
    wh.row_dimensions[1].height = 32
    wh.freeze_panes = "A2"

    # ---- Metrics
    wm = wb.create_sheet("Metrics")
    wm["A1"], wm["A1"].font = f"Validation metrics - {args.site} - run {args.run}", Font(name="Arial", size=12, bold=True)
    wm["A2"], wm["A2"].font = "Blue/yellow cells are inputs; everything else is a formula over Per_15min (only rows with Include = 1).", Font(name="Arial", size=9, italic=True)
    params = [(3, "Bin length (seconds)", 900, "GT is counted in 15-minute bins"),
              (4, "Min coverage of a bin to include it", 0.98, "bins processed for < 98% of their length (e.g. a smoke test) are excluded"),
              (5, "Flag a bin if |AI - GT| bicyclists exceeds", 2, "tolerance for the CHECK flag; a judgement call, edit freely")]
    for r_, lab, val, note in params:
        wm.cell(row=r_, column=1, value=lab).font = ARIAL
        c = wm.cell(row=r_, column=2, value=val)
        c.font, c.fill = INPUT_FONT, INPUT_FILL
        wm.cell(row=r_, column=3, value=note).font = Font(name="Arial", size=9, italic=True)
    J, K, Q, S, T, Rr = R("J"), R("K"), R("Q"), R("S"), R("T"), R("R")
    bic = [
        ("BICYCLISTS (bike lane + sidewalk, per 15-min bin)", None, None),
        ("Bins evaluated", f"=SUM({J})", "0"),
        ("Total AI", f"=SUMPRODUCT({J},{K})", "0"),
        ("Total GT", f"=SUMPRODUCT({J},{Q})", "0"),
        ("Bias: mean (AI - GT) per bin  (+ = overcount)", '=IF(B8=0,"",SUMPRODUCT(' + J + ',' + Rr + ')/B8)', "0.00"),
        ("MAE: mean |AI - GT| per bin", '=IF(B8=0,"",SUMPRODUCT(' + J + ',' + S + ')/B8)', "0.00"),
        ("RMSE per bin", '=IF(B8=0,"",SQRT(SUMPRODUCT(' + J + ',' + Rr + ',' + Rr + ')/B8))', "0.00"),
        ("WAPE: sum|err| / sum GT", '=IF(B10=0,"",SUMPRODUCT(' + J + ',' + S + ')/B10)', "0.0%"),
        ("Share of bins within +-1", '=IF(B8=0,"",SUMPRODUCT(' + J + ',--(' + S + '<=1))/B8)', "0.0%"),
        ("Share of bins within +-2", '=IF(B8=0,"",SUMPRODUCT(' + J + ',--(' + S + '<=2))/B8)', "0.0%"),
        ("Pearson r (AI vs GT per bin)",
         f'=IFERROR((B8*SUMPRODUCT({J},{K},{Q})-B9*B10)/SQRT((B8*SUMPRODUCT({J},{K},{K})-B9^2)*(B8*SUMPRODUCT({J},{Q},{Q})-B10^2)),"n/a")', "0.00"),
        ("Count-matched precision (approx.) = sum MIN(AI,GT) / sum AI", f'=IF(B9=0,"",SUMPRODUCT({J},{T})/B9)', "0.0%"),
        ("Count-matched recall (approx.) = sum MIN(AI,GT) / sum GT", f'=IF(B10=0,"",SUMPRODUCT({J},{T})/B10)', "0.0%"),
        ("Count-matched F1 (approx.)", '=IFERROR(2*B18*B19/(B18+B19),"")', "0.0%"),
    ]
    r_ = 7
    for lab, f, fmt in bic:
        wm.cell(row=r_, column=1, value=lab).font = BOLD if f is None else ARIAL
        if f is not None:
            c = wm.cell(row=r_, column=2, value=f)
            c.font, c.number_format = ARIAL, fmt
        r_ += 1
    # r_ is now 21
    wm.cell(row=r_, column=1, value="Note: there is no per-person ground truth, so precision/recall are count-based approximations "
                                  "(a missed rider in one bin cancelling a false rider in the same bin is not detected). Use the annotated clips for true error types.").font = Font(name="Arial", size=9, italic=True)
    r_ += 2
    V, W, X = R("V"), R("W"), R("X")
    ped = [
        ("PEDESTRIANS (North channels; reported, not the main target)", None, None),
        ("Total AI", f"=SUMPRODUCT({J},{V})", "0"),
        ("Total GT", f"=SUMPRODUCT({J},{W})", "0"),
        ("Bias: mean (AI - GT) per bin", f'=IF(B8=0,"",SUMPRODUCT({J},{X})/B8)', "0.00"),
        ("MAE per bin", f'=IF(B8=0,"",SUMPRODUCT({J},ABS({X}))/B8)', "0.00"),
        ("BIKE LANE vs SIDEWALK split (needs a zones file; 0 AI = zones not set)", None, None),
        ("AI bicyclists bike lane", f"=SUMPRODUCT({J},{R('L')})", "0"),
        ("GT bicyclists bike lane", f"=SUMPRODUCT({J},{R('O')})", "0"),
        ("AI bicyclists sidewalk", f"=SUMPRODUCT({J},{R('M')})", "0"),
        ("GT bicyclists sidewalk", f"=SUMPRODUCT({J},{R('P')})", "0"),
        ("AI bicyclists unassigned (no zone)", f"=SUMPRODUCT({J},{R('N')})", "0"),
    ]
    for lab, f, fmt in ped:
        wm.cell(row=r_, column=1, value=lab).font = BOLD if f is None else ARIAL
        if f is not None:
            c = wm.cell(row=r_, column=2, value=f)
            c.font, c.number_format = ARIAL, fmt
        r_ += 1
    wm.column_dimensions["A"].width = 70
    wm.column_dimensions["B"].width = 14
    wm.column_dimensions["C"].width = 70

    # ---- Run_Log
    wl = wb.create_sheet("Run_Log")
    wl.column_dimensions["A"].width = 36
    wl.column_dimensions["B"].width = 120
    r_ = 1
    for title, items in runlog_rows:
        wl.cell(row=r_, column=1, value=title).font = Font(name="Arial", size=11, bold=True)
        r_ += 1
        for k, v in items:
            wl.cell(row=r_, column=1, value=k).font = ARIAL
            c = wl.cell(row=r_, column=2, value=v)
            c.font, c.alignment = ARIAL, Alignment(wrap_text=True, vertical="top")
            r_ += 1
        r_ += 1
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)


def try_recalc(path):
    """Best effort: have LibreOffice (if present on the node) store cached values so viewers that don't calculate still show numbers."""
    so = shutil.which("soffice") or shutil.which("libreoffice")
    if not so:
        return False
    tmpdir = os.path.join(os.path.dirname(path), ".recalc_tmp")
    os.makedirs(tmpdir, exist_ok=True)
    r = subprocess.run([so, "--headless", "--convert-to", "xlsx", "--outdir", tmpdir, path], capture_output=True, timeout=180)
    out = os.path.join(tmpdir, os.path.basename(path))
    ok = r.returncode == 0 and os.path.exists(out)
    if ok:
        os.replace(out, path)
    shutil.rmtree(tmpdir, ignore_errors=True)
    return ok


# ----------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True, help="test_colonial folder")
    ap.add_argument("--run", required=True)
    ap.add_argument("--xls", default=None, help="GT .xls (default: <root>/ground_truth/Colonial_North_676806_07-10-2019.xls)")
    ap.add_argument("--gt-csv", default=None, help="ground_truth_bicycle_counts_v2.csv, for the cross-check")
    ap.add_argument("--mapping-csv", default=None, help="video_site_mapping.csv, for the cross-check")
    ap.add_argument("--video-dir", default=None, help="folder with the .NNN.mp4 chunks, to count them")
    ap.add_argument("--site", default="Colonial North")
    ap.add_argument("--location", default="Colonial - BL Colonial camera (SCU8OG)")
    ap.add_argument("--session-start", default="2019-07-10 00:00:00")
    ap.add_argument("--expect-chunks", default="1-48", help="chunks you intended to run, e.g. 1-24 for block 1")
    ap.add_argument("--ped-cols", default="1,2,3,4", help="Pedestrians-sheet columns for this camera (North = 1-4)")
    args = ap.parse_args()

    root = args.root
    xls = args.xls or os.path.join(root, "ground_truth", "Colonial_North_676806_07-10-2019.xls")
    sess = dt.datetime.strptime(args.session_start, "%Y-%m-%d %H:%M:%S")
    gt, gt_checks = parse_gt(xls, [int(c) for c in args.ped_cols.split(",")])

    bin_files = latest_per_chunk(os.path.join(root, "counts", args.run, "*_bin_counts*.csv"))
    log_files = latest_per_chunk(os.path.join(root, "logs", args.run, "*_run_log*.json"))
    if not bin_files:
        sys.exit(f"ERROR: no *_bin_counts.csv under {root}/counts/{args.run}/")
    expect = parse_chunks(args.expect_chunks)
    have = sorted(bin_files)
    missing = [c for c in expect if c not in bin_files]

    runlogs = {n: json.load(open(p)) for n, p in log_files.items()}
    clip_idx = {}
    for n, rl in runlogs.items():
        for cpath in rl.get("annotated_clips", []):
            m = re.search(r"_(\d{8})_(\d{4})-(\d{4})\.mp4$", cpath)
            if m:
                clip_idx[(m.group(1), m.group(2))] = os.path.relpath(cpath, root)

    rows, hours = [], []
    for n in have:
        with open(bin_files[n], newline="") as f:
            recs = list(csv.DictReader(f))
        for r in recs:
            b = int(r["bin_index"])
            t0 = dt.datetime.fromisoformat(r["bin_start"])
            t1 = dt.datetime.fromisoformat(r["bin_end"])
            g = gt.get(b, {"bic_bl": 0, "bic_sw": 0, "ped": 0})
            if b not in gt:
                print(f"WARNING: bin {b} has no ground truth row")
            rows.append({"video": r["video_file"], "clip": clip_idx.get((t0.strftime("%Y%m%d"), t0.strftime("%H%M")), ""),
                         "date": t0.strftime("%Y-%m-%d"), "t0": t0.strftime("%I:%M %p"), "t1": t1.strftime("%I:%M %p"),
                         "chunk": n, "secs": float(r["seconds_covered"]), "bin": b,
                         "ai_bic": int(r["bicyclists_total"]), "ai_bl": int(r["bicyclists_bike_lane"]),
                         "ai_sw": int(r["bicyclists_sidewalk"]), "ai_na": int(r["bicyclists_unassigned"]),
                         "gt_bl": g["bic_bl"], "gt_sw": g["bic_sw"], "ai_ped": int(r["pedestrians"]),
                         "gt_ped": g["ped"], "ai_veh": int(r["vehicles"])})
        hours.append({"chunk": n, "video": recs[0]["video_file"],
                      "t0": (sess + dt.timedelta(hours=n - 1)).strftime("%Y-%m-%d %I:%M %p")})

    # ---- python-side numbers for the text log (the workbook recomputes them with formulas)
    inc = [r["secs"] >= 0.98 * 900 for r in rows]
    m_bic = py_metrics([r["ai_bic"] for r in rows], [r["gt_bl"] + r["gt_sw"] for r in rows], inc)
    m_ped = py_metrics([r["ai_ped"] for r in rows], [r["gt_ped"] for r in rows], inc)
    n_partial = sum(not i for i in inc)

    # ---- checks
    checks = list(gt_checks)
    tot_xls = sum(v["bic_bl"] + v["bic_sw"] for v in gt.values())
    tot_csv = gt_csv_total(args.gt_csv, args.site)
    first_date = sess.strftime("%Y-%m-%d")
    mp = mapping_row(args.mapping_csv, args.site, first_date)
    checks.append(f"GT bicyclists, all {len(gt)} bins of the .xls (48 h): {tot_xls}")
    if tot_csv is not None:
        checks.append(f"GT CSV total for {args.site}: {tot_csv} -> {'MATCHES .xls' if tot_csv == tot_xls else 'MISMATCH vs .xls'}")
    if mp:
        checks.append(f"video_site_mapping: folder={mp['video_folder']}, gt_cyclist_count={mp['gt_cyclist_count']}, "
                      f"n_mp4_chunks={mp['n_mp4_chunks']} -> GT is a {tot_xls}-count over {len(gt) // 4} h of bins; "
                      f"chunk count {'consistent' if int(mp['n_mp4_chunks']) * 4 == len(gt) else 'NOT consistent'} with the {len(gt)} GT bins")
    if args.video_dir and os.path.isdir(args.video_dir):
        nmp4 = len(glob.glob(os.path.join(args.video_dir, "*.mp4")))
        checks.append(f"mp4 files found in video folder: {nmp4}")

    chunk_rows, warn = [], []
    prev_end = None
    for n in have:
        rl = runlogs.get(n)
        if not rl:
            warn.append(f"chunk {n}: no run log found")
            continue
        dur = rl["duration_s_container"]
        if abs(dur - 3600) > 5:
            warn.append(f"chunk {n}: duration {dur:.1f} s differs from 3600 s by more than 5 s (gap/overlap?)")
        if rl["partial_run"]:
            warn.append(f"chunk {n}: PARTIAL run (start-frame/max-frames used), {rl['duration_s_processed']:.0f} s processed")
        if rl["frames_processed"] != rl["frames_in_container_header"] and not rl["partial_run"]:
            warn.append(f"chunk {n}: processed {rl['frames_processed']} frames but container header says {rl['frames_in_container_header']}")
        start = dt.datetime.fromisoformat(rl["clip_start_clock"])
        if prev_end is not None and abs((start - prev_end).total_seconds()) > 5:
            warn.append(f"chunk {n}: starts {start} but previous processed chunk ends {prev_end} (non-adjacent chunks)")
        prev_end = start + dt.timedelta(seconds=dur)
        chunk_rows.append((n, rl))
    if missing:
        warn.append(f"expected chunks not found for this run: {missing}")

    # ---- Run_Log sheet content
    glob_items = [
        ("Run name", args.run), ("Site / location", f"{args.site} - {args.location}"),
        ("Video session start (camera clock)", f"{args.session_start} (chunk .001); verified from on-screen clock on .001 = 2019-07-10 12:00:00 AM, .008 = 7:00:01 AM, .025 = 2019-07-11 12:00:01 AM"),
        ("Chunks processed", f"{len(have)}: {have}"),
        ("Chunks expected", f"{len(expect)} ({args.expect_chunks}); missing: {missing if missing else 'none'}"),
        ("Videos merged", "0 - each hourly mp4 is processed on its own and the results are joined by clock time; "
                          "a rider crossing an hour boundary is split across two chunks (rare)"),
        ("Total video duration processed", f"{sum(rl['duration_s_processed'] for _, rl in chunk_rows) / 3600:.2f} h"),
        ("15-min bins in workbook", f"{len(rows)} ({len(rows) - n_partial} evaluated, {n_partial} partial/excluded)"),
        ("Ground truth file", xls), ("Weights", chunk_rows[0][1]["weights"] if chunk_rows else "?"),
        ("Total runtime (sum over chunks)", f"{sum(rl['runtime_s'] for _, rl in chunk_rows) / 60:.1f} min"),
    ]
    runlog_rows = [("RUN SUMMARY", glob_items), ("CHECKS", [(f"check {i + 1}", c) for i, c in enumerate(checks)]),
                   ("WARNINGS", [(f"warning {i + 1}", w) for i, w in enumerate(warn)] or [("none", "")])]
    for n, rl in chunk_rows:
        p = rl["params"]
        runlog_rows.append((f"CHUNK {n:03d}  {rl['video_file']}", [
            ("clip start (clock)", rl["clip_start_clock"]),
            ("duration container / processed (s)", f"{rl['duration_s_container']} / {rl['duration_s_processed']}"),
            ("frames processed / header", f"{rl['frames_processed']} / {rl['frames_in_container_header']}  (fps {rl['fps']}, {rl['resolution'][0]}x{rl['resolution'][1]})"),
            ("final counts (AI)", f"bicyclists {rl['final_counts']['bicyclists']}, pedestrians {rl['final_counts']['pedestrians']}, vehicles {rl['final_counts']['vehicles']}"),
            ("params", f"imgsz {p.get('imgsz')}, det_conf {p.get('det_conf')}, track_thr {p.get('track_thr')}, person_min_sec {p.get('person_min_sec')}, stride {p.get('stride')}, zones {p.get('zones_json')}"),
            ("runtime / device", f"{rl['runtime_s']} s / {rl['device']}"),
            ("annotated clips", ", ".join(os.path.relpath(c, root) for c in rl["annotated_clips"])),
            ("finished", rl["finished"])]))

    outdir = os.path.join(root, "validation")
    os.makedirs(outdir, exist_ok=True)
    stem = f"colonial_north_{args.run}_validation"
    xlsx_path = unique_path(os.path.join(outdir, stem + ".xlsx"))
    notes = []
    build_workbook(rows, hours, runlog_rows, notes, args, xlsx_path)
    recalc_ok = try_recalc(xlsx_path)

    # ---- text log
    log_path = unique_path(os.path.join(outdir, stem + "_log.txt"))
    L = []
    L.append(f"VALIDATION LOG  {args.site}  run={args.run}  written {dt.datetime.now():%Y-%m-%d %H:%M:%S}")
    L.append("=" * 100)
    for t, items in runlog_rows[:3]:
        L.append(f"\n[{t}]")
        L += [f"  {k}: {v}" if v != "" else f"  {k}" for k, v in items]
    for title, m in (("BICYCLISTS (bike lane + sidewalk) per 15-min bin", m_bic), ("PEDESTRIANS per 15-min bin", m_ped)):
        L.append(f"\n[{title}]  (bins with full coverage only)")
        for k, v in m.items():
            L.append(f"  {k}: {v:.3f}" if isinstance(v, float) else f"  {k}: {v}")
    L.append("\n[PER-HOUR]")
    L.append("  chunk  hour start            AI bic  GT bic   AI ped  GT ped")
    for n in have:
        rr = [r for r in rows if r["chunk"] == n and r["secs"] >= 0.98 * 900]
        L.append(f"  {n:5d}  {(sess + dt.timedelta(hours=n - 1)):%Y-%m-%d %I:%M %p}   {sum(r['ai_bic'] for r in rr):6d}  {sum(r['gt_bl'] + r['gt_sw'] for r in rr):6d}   "
                 f"{sum(r['ai_ped'] for r in rr):6d}  {sum(r['gt_ped'] for r in rr):6d}")
    L.append("\n[BINS FLAGGED |AI-GT| > 2 bicyclists]")
    flagged = [r for r in rows if r["secs"] >= 0.98 * 900 and abs(r["ai_bic"] - r["gt_bl"] - r["gt_sw"]) > 2]
    L += [f"  {r['date']} {r['t0']}-{r['t1']}  AI {r['ai_bic']}  GT {r['gt_bl'] + r['gt_sw']}   clip: {r['clip']}" for r in flagged] or ["  none"]
    L.append(f"\nWorkbook: {xlsx_path}  (cached values stored via LibreOffice: {recalc_ok}; Excel recalculates on open either way)")
    with open(log_path, "w") as f:
        f.write("\n".join(L) + "\n")
    print("\n".join(L))
    print(f"\nWrote:\n  {xlsx_path}\n  {log_path}")


if __name__ == "__main__":
    main()
